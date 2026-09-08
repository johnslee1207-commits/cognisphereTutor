"""Local course artifact exporter for OpenMAIC-style stage documents."""

from __future__ import annotations

from dataclasses import dataclass
from html import escape
import json
from pathlib import Path
from typing import Any, Mapping
from zipfile import ZIP_DEFLATED, ZipFile

from pptx import Presentation
from pptx.util import Inches


@dataclass(frozen=True)
class ExportedCourseArtifact:
    artifact_id: str
    artifact_type: str
    format: str
    path: Path
    content_type: str
    size_bytes: int


def export_course_artifact(
    *,
    artifact_id: str,
    artifact_type: str,
    format: str,
    course_id: str,
    version: str,
    stage: Mapping[str, Any],
    scenes: list[Mapping[str, Any]],
    output_root: Path | None = None,
) -> ExportedCourseArtifact:
    """Write a real local HTML, PPTX, or MAIC-ZIP artifact."""
    if output_root is None:
        from cognispheretutor.services.path_service import get_path_service

        root = get_path_service().get_public_outputs_root() / "course_exports"
    else:
        root = output_root
    artifact_dir = root / _safe_name(course_id) / _safe_name(version) / artifact_id
    artifact_dir.mkdir(parents=True, exist_ok=True)

    if format == "html":
        path = artifact_dir / "course.html"
        path.write_text(_render_html(stage, scenes), encoding="utf-8")
        content_type = "text/html; charset=utf-8"
    elif format == "pptx":
        path = artifact_dir / "course.pptx"
        _write_pptx(path, stage, scenes)
        content_type = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    elif format == "maic-zip":
        path = artifact_dir / "course.maic.zip"
        _write_maic_zip(path, stage, scenes)
        content_type = "application/zip"
    else:
        raise ValueError(f"unsupported course export format: {format}")

    return ExportedCourseArtifact(
        artifact_id=artifact_id,
        artifact_type=artifact_type,
        format=format,
        path=path,
        content_type=content_type,
        size_bytes=path.stat().st_size,
    )


def _render_html(stage: Mapping[str, Any], scenes: list[Mapping[str, Any]]) -> str:
    title = escape(str(stage.get("name") or stage.get("id") or "Course"))
    description = escape(str(stage.get("description") or ""))
    sections = "\n".join(_render_scene_html(scene) for scene in scenes)
    return f"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{title}</title>
  <style>
    body {{ font-family: Arial, sans-serif; margin: 32px; color: #1f2937; }}
    main {{ max-width: 920px; margin: 0 auto; }}
    section {{ border-top: 1px solid #d1d5db; padding: 24px 0; }}
    h1, h2 {{ color: #111827; }}
    .kind {{ color: #4b5563; font-size: 13px; text-transform: uppercase; }}
    li {{ margin: 6px 0; }}
  </style>
</head>
<body>
<main>
  <h1>{title}</h1>
  <p>{description}</p>
  {sections}
</main>
</body>
</html>
"""


def _render_scene_html(scene: Mapping[str, Any]) -> str:
    title = escape(str(scene.get("title") or scene.get("id") or "Scene"))
    scene_type = escape(str(scene.get("type") or "scene"))
    content = scene.get("content") if isinstance(scene.get("content"), Mapping) else {}
    bullets = _content_bullets(content)
    items = "\n".join(f"<li>{escape(item)}</li>" for item in bullets)
    return f"""<section>
  <div class="kind">{scene_type}</div>
  <h2>{title}</h2>
  <ul>{items}</ul>
</section>"""


def _write_pptx(path: Path, stage: Mapping[str, Any], scenes: list[Mapping[str, Any]]) -> None:
    prs = Presentation()
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = str(stage.get("name") or stage.get("id") or "Course")
    title_slide.placeholders[1].text = str(stage.get("description") or "")

    for scene in scenes:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = str(scene.get("title") or scene.get("id") or "Scene")
        body = slide.placeholders[1].text_frame
        body.clear()
        content = scene.get("content") if isinstance(scene.get("content"), Mapping) else {}
        for index, item in enumerate(_content_bullets(content)):
            paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
            paragraph.text = item
            paragraph.level = 0

    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.save(path)


def _write_maic_zip(path: Path, stage: Mapping[str, Any], scenes: list[Mapping[str, Any]]) -> None:
    manifest = {
        "formatVersion": 1,
        "stage": dict(stage),
        "scenes": [dict(scene) for scene in scenes],
        "mediaIndex": {},
    }
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2))
        archive.writestr("course.html", _render_html(stage, scenes))


def _content_bullets(content: Mapping[str, Any]) -> list[str]:
    if content.get("type") == "quiz":
        questions = content.get("questions") if isinstance(content.get("questions"), list) else []
        return [str(q.get("prompt") or q) for q in questions if isinstance(q, Mapping)] or [
            "Answer the checkpoint question and review evidence."
        ]
    if content.get("type") == "pbl":
        project = content.get("project") if isinstance(content.get("project"), Mapping) else {}
        tasks = project.get("tasks") if isinstance(project.get("tasks"), list) else []
        return [str(task) for task in tasks] or ["Complete the project task and submit evidence."]
    if content.get("type") == "interactive":
        return ["Explore the embedded interactive activity.", "Record observations as evidence."]
    canvas = content.get("canvas") if isinstance(content.get("canvas"), Mapping) else {}
    elements = canvas.get("elements") if isinstance(canvas.get("elements"), list) else []
    texts = [
        str(item.get("text") or item.get("content"))
        for item in elements
        if isinstance(item, Mapping) and (item.get("text") or item.get("content"))
    ]
    return texts or ["Review the concept and connect it to the learning objective."]


def _safe_name(value: str) -> str:
    cleaned = "".join(ch if ch.isalnum() or ch in {"-", "_", "."} else "_" for ch in value)
    return cleaned.strip("._") or "course"
